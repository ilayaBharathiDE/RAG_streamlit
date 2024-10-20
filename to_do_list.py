import streamlit as st
import json
from datetime import datetime
import openpyxl
from openpyxl.styles import Font
import os

# Function to load tasks from a JSON file
def load_tasks():
    try:
        with open("tasks.json", "r") as f:
            tasks = json.load(f)
            # Add status to existing tasks if missing
            for task in tasks:
                if "status" not in task:
                    task["status"] = "To Do"
                if "status_change_date" not in task:
                    task["status_change_date"] = task.get("date", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
            return tasks
    except FileNotFoundError:
        return []

# Function to save tasks to a JSON file
def save_tasks(tasks):
    with open("tasks.json", "w") as f:
        json.dump(tasks, f)

# Function to update Excel history
def update_excel_history(tasks):
    excel_file = "todo_history.xlsx"
    
    if os.path.exists(excel_file):
        wb = openpyxl.load_workbook(excel_file)
    else:
        wb = openpyxl.Workbook()
        
    sheet = wb.active
    sheet.title = "To-Do History"
    
    # Clear existing content
    sheet.delete_rows(1, sheet.max_row)
    
    # Write headers
    headers = ["Task", "Status", "Created Date", "Last Status Change"]
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)
    
    # Write tasks
    for row, task in enumerate(tasks, start=2):
        sheet.cell(row=row, column=1).value = task['task']
        sheet.cell(row=row, column=2).value = task['status']
        sheet.cell(row=row, column=3).value = task['date']
        sheet.cell(row=row, column=4).value = task['status_change_date']
    
    # Save the workbook
    wb.save(excel_file)

# Custom CSS to add more color and style
st.markdown("""
<style>
    body {
        color: #333333;
    }
    .stApp {
        background: linear-gradient(135deg, #f5f7fa 0%, #c3cfe2 100%);
    }
    .stButton>button {
        background-color: #00bcd4;
        color: white;
        width: 100%;
        height: 40px;
    }
    .stTextInput>div>div>input {
        background-color: #ffffff;
    }
    h1 {
        color: #006064;
    }
    h2 {
        color: #00838f;
    }
    .task-card {
        background-color: rgba(255, 255, 255, 0.8);
        border-radius: 5px;
        padding: 10px;
        margin-bottom: 10px;
        box-shadow: 0 2px 5px rgba(0,0,0,0.1);
        overflow-wrap: break-word;
        word-wrap: break-word;
        hyphens: auto;
    }
    .horizontal-stats {
        display: flex;
        justify-content: space-between;
        align-items: center;
        background-color: rgba(255, 255, 255, 0.8);
        padding: 10px;
        border-radius: 5px;
        margin-bottom: 20px;
    }
    .stat-item {
        text-align: center;
    }
    @keyframes flow {
        0% {
            background-position: 0% 50%;
        }
        50% {
            background-position: 100% 50%;
        }
        100% {
            background-position: 0% 50%;
        }
    }
    .flow-background {
        background: linear-gradient(270deg, #f5f7fa, #c3cfe2, #e0eafc, #cfdef3);
        background-size: 400% 400%;
        animation: flow 15s ease infinite;
        position: fixed;
        top: 0;
        left: 0;
        right: 0;
        bottom: 0;
        z-index: -1;
    }
</style>
<div class="flow-background"></div>
""", unsafe_allow_html=True)

# Initialize the app
st.title("Enhanced Daily To-Do List")

# Load existing tasks
tasks = load_tasks()

# Display task statistics
total_tasks = len(tasks)
todo_tasks = sum(1 for task in tasks if task.get("status", "To Do") == "To Do")
inprogress_tasks = sum(1 for task in tasks if task.get("status", "To Do") == "In Progress")
done_tasks = sum(1 for task in tasks if task.get("status", "To Do") == "Done")

st.markdown("""
<div class="horizontal-stats">
    <div class="stat-item">
        <h3>Total tasks</h3>
        <p>{}</p>
    </div>
    <div class="stat-item">
        <h3>To Do</h3>
        <p>{}</p>
    </div>
    <div class="stat-item">
        <h3>In Progress</h3>
        <p>{}</p>
    </div>
    <div class="stat-item">
        <h3>Done</h3>
        <p>{}</p>
    </div>
</div>
""".format(total_tasks, todo_tasks, inprogress_tasks, done_tasks), unsafe_allow_html=True)

# Input for adding a new task
new_task = st.text_input("Add a new task")
task_status = st.selectbox("Task Status", ["To Do", "In Progress", "Done"])
if st.button("Add Task"):
    if new_task:
        current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        tasks.append({
            "task": new_task,
            "status": task_status,
            "date": current_time,
            "status_change_date": current_time
        })
        save_tasks(tasks)
        update_excel_history(tasks)
        st.success("Task added successfully!")
        st.rerun()
    else:
        st.warning("Please enter a task.")

# Display and manage existing tasks
for status in ["To Do", "In Progress", "Done"]:
    st.subheader(status)
    for index, task in enumerate(tasks):
        if task.get("status", "To Do") == status:
            with st.container():
                col1, col2, col3, col4 = st.columns([3, 2, 2, 1])
                with col1:
                    st.markdown(f"<div class='task-card'>{task['task']}</div>", unsafe_allow_html=True)
                with col2:
                    new_status = st.selectbox("", ["To Do", "In Progress", "Done"], index=["To Do", "In Progress", "Done"].index(task.get("status", "To Do")), key=f"status_{index}")
                    if new_status != task.get("status", "To Do"):
                        tasks[index]["status"] = new_status
                        tasks[index]["status_change_date"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                        save_tasks(tasks)
                        update_excel_history(tasks)
                        st.rerun()
                with col3:
                    st.write(f"Created: {task.get('date', 'No date')}")
                    st.write(f"Status changed: {task.get('status_change_date', 'No date')}")
                with col4:
                    if st.button("Remove", key=f"remove_{index}"):
                        tasks.pop(index)
                        save_tasks(tasks)
                        update_excel_history(tasks)
                        st.rerun()

# Add a button to clear all completed tasks
if st.button("Clear Completed Tasks"):
    tasks = [task for task in tasks if task.get("status", "To Do") != "Done"]
    save_tasks(tasks)
    update_excel_history(tasks)
    st.success("Completed tasks cleared!")
    st.rerun()

# Add a button to download the Excel history
if st.button("Download Task History"):
    update_excel_history(tasks)
    with open("todo_history.xlsx", "rb") as file:
        btn = st.download_button(
            label="Download Excel",
            data=file,
            file_name="todo_history.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
